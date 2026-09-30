"""
to_pptx.py — PEMBUAT LAPORAN POWERPOINT
Membaca hasil data/output/rekap.xlsx lalu membuat laporan.pptx berisi:
  - Slide judul
  - Slide ringkasan (kartu angka besar)
  - Grafik pemasukan vs pengeluaran per wilayah
  - Grafik komposisi pengeluaran per kategori
  - Slide penutup

Cara pakai:  python scripts/to_pptx.py
Jalankan SETELAH process.py.
"""
import os
import sys
import pandas as pd
import yaml

import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN

BASE = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))

# Palet warna korporat (bisa disesuaikan)
BIRU_TUA = RGBColor(0x1F, 0x3A, 0x5F)
BIRU = RGBColor(0x2E, 0x6C, 0xB5)
HIJAU = RGBColor(0x2E, 0xA0, 0x6C)
MERAH = RGBColor(0xC0, 0x39, 0x2B)
ABU = RGBColor(0x6B, 0x72, 0x80)
PUTIH = RGBColor(0xFF, 0xFF, 0xFF)
HEX_BIRU = "#2E6CB5"
HEX_HIJAU = "#2EA06C"
HEX_MERAH = "#C0392B"
PALET = ["#2E6CB5", "#2EA06C", "#E67E22", "#8E44AD", "#C0392B", "#16A085", "#7F8C8D"]


def muat_config():
    with open(os.path.join(BASE, "config.yaml"), "r", encoding="utf-8") as f:
        return yaml.safe_load(f)


def path_absolut(p):
    return p if os.path.isabs(p) else os.path.join(BASE, p)


def fmt_rp(x, mu="Rp"):
    return f"{mu} {x:,.0f}"


def grafik_per_grup(per_grup, grup, tmp_dir, mu):
    plt.figure(figsize=(8, 4.2))
    x = range(len(per_grup))
    lebar = 0.38
    plt.bar([i - lebar/2 for i in x], per_grup["Pemasukan"], lebar,
            label="Pemasukan", color=HEX_HIJAU)
    plt.bar([i + lebar/2 for i in x], per_grup["Pengeluaran"], lebar,
            label="Pengeluaran", color=HEX_MERAH)
    plt.xticks(list(x), per_grup[grup], rotation=0)
    plt.ylabel(f"Nilai ({mu})")
    plt.title(f"Pemasukan vs Pengeluaran per {grup.capitalize()}")
    plt.legend()
    plt.gca().yaxis.set_major_formatter(
        matplotlib.ticker.FuncFormatter(lambda v, _: f"{v/1e6:.0f} jt"))
    plt.tight_layout()
    path = os.path.join(tmp_dir, "grafik_grup.png")
    plt.savefig(path, dpi=140)
    plt.close()
    return path


def grafik_kategori(per_kat, tmp_dir):
    data = per_kat[per_kat["Pengeluaran"] > 0]
    if data.empty:
        return None
    plt.figure(figsize=(6.5, 5))
    plt.pie(data["Pengeluaran"], labels=data["kategori"], autopct="%1.0f%%",
            colors=PALET[:len(data)], startangle=90,
            textprops={"fontsize": 9})
    plt.title("Komposisi Pengeluaran per Kategori")
    plt.tight_layout()
    path = os.path.join(tmp_dir, "grafik_kategori.png")
    plt.savefig(path, dpi=140)
    plt.close()
    return path


def kotak_teks(slide, kiri, atas, lebar, tinggi, teks, ukuran, warna,
               bold=False, align=PP_ALIGN.LEFT):
    box = slide.shapes.add_textbox(kiri, atas, lebar, tinggi)
    tf = box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = teks
    run.font.size = Pt(ukuran)
    run.font.bold = bold
    run.font.color.rgb = warna
    return box


def kartu(slide, kiri, atas, lebar, tinggi, judul, nilai, warna):
    kotak = slide.shapes.add_shape(1, kiri, atas, lebar, tinggi)  # rectangle
    kotak.fill.solid()
    kotak.fill.fore_color.rgb = warna
    kotak.line.fill.background()
    tf = kotak.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    r = p.add_run(); r.text = judul
    r.font.size = Pt(13); r.font.color.rgb = PUTIH
    p2 = tf.add_paragraph(); p2.alignment = PP_ALIGN.CENTER
    r2 = p2.add_run(); r2.text = nilai
    r2.font.size = Pt(20); r2.font.bold = True; r2.font.color.rgb = PUTIH


def main():
    cfg = muat_config()
    mu = cfg["ppt"]["mata_uang"]
    out_dir = path_absolut(cfg["folders"]["output"])
    rekap_path = os.path.join(out_dir, "rekap.xlsx")
    if not os.path.exists(rekap_path):
        print("[!] rekap.xlsx belum ada. Jalankan dulu: python scripts/process.py")
        sys.exit(1)

    tmp_dir = os.path.join(out_dir, "_grafik")
    os.makedirs(tmp_dir, exist_ok=True)

    xls = pd.ExcelFile(rekap_path)
    ringkas = pd.read_excel(xls, "Ringkasan")
    nama_grup_sheet = [s for s in xls.sheet_names if s.startswith("Per ") and s != "Per Kategori"][0]
    per_grup = pd.read_excel(xls, nama_grup_sheet)
    grup = per_grup.columns[0]
    per_kat = pd.read_excel(xls, "Per Kategori")

    nilai = dict(zip(ringkas["Keterangan"], ringkas["Nilai"]))
    total_masuk = nilai.get("Total Pemasukan (Kredit)", 0)
    total_keluar = nilai.get("Total Pengeluaran (Debit)", 0)
    net = total_masuk - total_keluar
    jml_trx = int(nilai.get("Jumlah Transaksi", 0))

    # buat grafik
    g1 = grafik_per_grup(per_grup, grup, tmp_dir, mu)
    g2 = grafik_kategori(per_kat, tmp_dir)

    # bangun presentasi
    template = cfg["ppt"].get("template") or ""
    prs = Presentation(path_absolut(template)) if template and os.path.exists(path_absolut(template)) else Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank = prs.slide_layouts[6]
    W = prs.slide_width

    # --- Slide 1: Judul ---
    s = prs.slides.add_slide(blank)
    bg = s.shapes.add_shape(1, 0, 0, W, prs.slide_height)
    bg.fill.solid(); bg.fill.fore_color.rgb = BIRU_TUA; bg.line.fill.background()
    kotak_teks(s, Inches(1), Inches(2.6), Inches(11.3), Inches(1.2),
               cfg["ppt"]["judul"], 40, PUTIH, bold=True, align=PP_ALIGN.CENTER)
    kotak_teks(s, Inches(1), Inches(3.8), Inches(11.3), Inches(0.8),
               cfg["ppt"]["subjudul"], 22, RGBColor(0xBB, 0xCC, 0xEE), align=PP_ALIGN.CENTER)
    kotak_teks(s, Inches(1), Inches(6.4), Inches(11.3), Inches(0.6),
               cfg["ppt"]["perusahaan"], 16, RGBColor(0x99, 0xAA, 0xCC), align=PP_ALIGN.CENTER)

    # --- Slide 2: Ringkasan kartu ---
    s = prs.slides.add_slide(blank)
    kotak_teks(s, Inches(0.6), Inches(0.4), Inches(12), Inches(0.8),
               "Ringkasan Eksekutif", 30, BIRU_TUA, bold=True)
    lebar_k = Inches(3.9); tinggi_k = Inches(1.7); atas_k = Inches(1.6); gap = Inches(0.25)
    kartu(s, Inches(0.6), atas_k, lebar_k, tinggi_k, "Total Pemasukan", fmt_rp(total_masuk, mu), HIJAU)
    kartu(s, Inches(0.6)+lebar_k+gap, atas_k, lebar_k, tinggi_k, "Total Pengeluaran", fmt_rp(total_keluar, mu), MERAH)
    kartu(s, Inches(0.6)+2*(lebar_k+gap), atas_k, lebar_k, tinggi_k, "Net / Selisih", fmt_rp(net, mu), BIRU)
    kotak_teks(s, Inches(0.6), Inches(3.8), Inches(12), Inches(0.6),
               f"Jumlah transaksi: {jml_trx:,}  |  Dikelompokkan per: {grup}", 16, ABU)
    # tabel ringkas per grup
    baris = "\n".join(
        f"  • {r[grup]}:  masuk {fmt_rp(r['Pemasukan'], mu)}  |  keluar {fmt_rp(r['Pengeluaran'], mu)}"
        for _, r in per_grup.iterrows())
    kotak_teks(s, Inches(0.6), Inches(4.5), Inches(12), Inches(2.5),
               f"Rincian per {grup.capitalize()}:\n{baris}", 14, RGBColor(0x33, 0x33, 0x33))

    # --- Slide 3: Grafik per grup ---
    s = prs.slides.add_slide(blank)
    kotak_teks(s, Inches(0.6), Inches(0.4), Inches(12), Inches(0.8),
               f"Pemasukan vs Pengeluaran per {grup.capitalize()}", 28, BIRU_TUA, bold=True)
    s.shapes.add_picture(g1, Inches(1.5), Inches(1.5), width=Inches(10.3))

    # --- Slide 4: Grafik kategori ---
    if g2:
        s = prs.slides.add_slide(blank)
        kotak_teks(s, Inches(0.6), Inches(0.4), Inches(12), Inches(0.8),
                   "Komposisi Pengeluaran per Kategori", 28, BIRU_TUA, bold=True)
        s.shapes.add_picture(g2, Inches(3.7), Inches(1.4), width=Inches(6))

    # --- Slide 5: Penutup ---
    s = prs.slides.add_slide(blank)
    bg = s.shapes.add_shape(1, 0, 0, W, prs.slide_height)
    bg.fill.solid(); bg.fill.fore_color.rgb = BIRU_TUA; bg.line.fill.background()
    kotak_teks(s, Inches(1), Inches(3), Inches(11.3), Inches(1),
               "Terima Kasih", 40, PUTIH, bold=True, align=PP_ALIGN.CENTER)
    kotak_teks(s, Inches(1), Inches(4.1), Inches(11.3), Inches(0.8),
               "Laporan dibuat otomatis oleh sistem.", 18,
               RGBColor(0xBB, 0xCC, 0xEE), align=PP_ALIGN.CENTER)

    ppt_path = os.path.join(out_dir, "laporan.pptx")
    prs.save(ppt_path)
    print(f"[OK] PPT tersimpan: {ppt_path}")
    print(f"     {len(prs.slides.__iter__.__self__._sldIdLst)} slide dibuat.")


if __name__ == "__main__":
    main()
