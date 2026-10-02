#!/usr/bin/env python3
"""Olah export absensi HRIS bulanan menjadi rekap penggajian.

Pemakaian:
    python absensi.py data/attendance-2026-10-01-to-2026-10-31-report.xlsx

Opsional (di folder data/):
    izin_cuti.csv   -> override hari "Absent" menjadi Izin/Sakit/Cuti
    karyawan.csv    -> gaji pokok per karyawan untuk hitung potongan/upah

Aturan bisnis ada di config.json.
"""
import argparse
import json
import sys
from pathlib import Path

import pandas as pd
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

BASE = Path(__file__).parent
BULAN = {"jan": 1, "feb": 2, "mar": 3, "apr": 4, "mei": 5, "may": 5, "jun": 6, "jul": 7,
         "agu": 8, "aug": 8, "agt": 8, "sep": 9, "okt": 10, "oct": 10, "nov": 11, "des": 12, "dec": 12}
KOLOM_WAJIB = ["Date", "Day", "Employee Name", "NIP", "Branch", "Post", "Schedule In",
               "Schedule Out", "Check In At", "Check Out At", "Attendance Code"]


# ---------- util waktu ----------
def tgl(s):
    d, b, y = str(s).split()
    return pd.Timestamp(int(y), BULAN[b.lower()[:3]], int(d))


def menit(s):
    """'HH:MM[:SS]' -> menit sejak 00:00 (pecahan untuk detik). None jika kosong/X."""
    s = str(s).strip()
    if len(s) < 5 or s[2] != ":":
        return None
    h, m = int(s[:2]), int(s[3:5])
    sec = int(s[6:8]) if len(s) >= 8 else 0
    return h * 60 + m + sec / 60


def hhmm(x):
    return f"{int(x) // 60:02d}:{int(x) % 60:02d}"


def valid(s):
    return str(s).strip() not in ("", "X", "nan", "None", "Off", "-")


# ---------- shift ----------
def siapkan_shift(cfg):
    out = []
    for s in cfg["shifts"]:
        out.append({"nama": s["nama"], "masuk": menit(s["masuk"]), "pulang": menit(s["pulang"]),
                    "cabang": s.get("cabang"), "hari": s.get("hari")})
    return out


def kandidat(shifts, cabang, hari):
    return [s for s in shifts if (not s["cabang"] or cabang in s["cabang"]) and (not s["hari"] or hari in s["hari"])]


def shift_terjadwal(row, cfg, shifts):
    """Shift sesuai jadwal HRIS; jadwal tak masuk akal dipetakan lewat alias."""
    si, so = row["Schedule In"], row["Schedule Out"]
    if not (valid(si) and valid(so)):
        return None
    key = f"{si}-{so}"
    if key in cfg["alias_jadwal"]:
        nama = cfg["alias_jadwal"][key]
        return next(s for s in shifts if s["nama"] == nama)
    a, b = menit(si), menit(so)
    if b <= a or b - a > 12 * 60:
        return None
    return {"nama": f"Jadwal {si}-{so}", "masuk": a, "pulang": b, "cabang": None, "hari": None}


def biaya(sh, ci, co):
    """Seberapa jauh check-in/out dari sebuah shift (menit). Makin kecil makin cocok."""
    c = abs(ci - sh["masuk"])
    if co is not None and co > ci:
        c += 0.5 * abs(co - sh["pulang"])
    return c


def tentukan_shift(row, cfg, shifts):
    """Return (shift_aktual, shift_jadwal, status, biaya)."""
    ci, co = menit(row["Check In At"]), menit(row["Check Out At"])
    jad = shift_terjadwal(row, cfg, shifts)
    if jad and row["Branch"] in cfg["cabang_shift_tetap"]:
        return jad, jad, "sesuai jadwal (shift tetap)", abs(ci - jad["masuk"])
    cands = kandidat(shifts, row["Branch"], row["Day"])
    if jad and all(c["nama"] != jad["nama"] for c in cands):
        cands.append(jad)
    best = min(cands, key=lambda s: biaya(s, ci, co))
    if jad is None:
        return best, None, "shift ditebak (jadwal HRIS tidak valid)", biaya(best, ci, co)
    b_jad, b_best = biaya(jad, ci, co), biaya(best, ci, co)
    if best["hari"] and best["nama"] != jad["nama"] and b_best < b_jad:
        return best, jad, "shift khusus hari", b_best
    if b_jad - b_best >= cfg["ganti_shift_selisih_menit"] and best["nama"] != jad["nama"]:
        return best, jad, "shift berubah (otomatis)", b_best
    return jad, jad, "sesuai jadwal", b_jad


# ---------- baca data ----------
def baca_export(path):
    df = pd.read_excel(path, dtype=str)
    kurang = [c for c in KOLOM_WAJIB if c not in df.columns]
    if kurang:
        sys.exit(f"Kolom export tidak ditemukan: {kurang}. Format export HRIS berubah?")
    df["tanggal"] = df["Date"].map(tgl)
    df["Employee Name"] = df["Employee Name"].str.strip()
    df["NIP"] = df["NIP"].fillna("").str.strip()
    return df


def baca_opsional(nama):
    p = BASE / "data" / nama
    return pd.read_csv(p, dtype=str).fillna("") if p.exists() else None


# ---------- hitung per hari ----------
def aturan_lembur(lembur, cabang, tanggal, nip):
    """Cari aturan lembur manual yang cocok (cabang '*' = semua, nip kosong = semua). Return tarif/jam atau None."""
    if lembur is None:
        return None
    for _, a in lembur.iterrows():
        if a["cabang"] not in ("*", cabang):
            continue
        if a.get("nip", "") not in ("", nip):
            continue
        awal = pd.Timestamp(a["tanggal_mulai"])
        akhir = pd.Timestamp(a["tanggal_selesai"]) if a.get("tanggal_selesai", "") else awal
        if awal <= pd.Timestamp(tanggal) <= akhir:
            return float(a["tarif_per_jam"] or 0)
    return None


def hitung_hari(row, cfg, shifts, lembur_rules=None):
    r = {"Tanggal": row["tanggal"].date(), "Hari": row["Day"], "Nama": row["Employee Name"], "NIP": row["NIP"],
         "Cabang": row["Branch"], "Posisi": row["Post"], "Jadwal HRIS": "", "Shift Aktual": "", "Status Shift": "",
         "Check In": "", "Check Out": "", "Status": row["Attendance Code"], "Telat (mnt)": 0,
         "Pulang Cepat (mnt)": 0, "Lembur (mnt)": 0, "Lembur Dibayar (jam)": 0.0, "Upah Lembur (Rp)": 0, "Potongan Telat (Rp)": 0, "Jam Kerja": 0.0,
         "Double Shift": "", "Telat HRIS (mnt)": "", "Catatan": ""}
    cat = []
    code = row["Attendance Code"]
    if valid(row["Schedule In"]):
        r["Jadwal HRIS"] = f"{row['Schedule In']}-{row['Schedule Out']}"
    if code != "Present" or not valid(row["Check In At"]):
        if code == "Present":
            r["Status"] = "Absent"
        return r, cat

    ci, co = menit(row["Check In At"]), menit(row["Check Out At"])
    r["Check In"], r["Check Out"] = row["Check In At"], row["Check Out At"] if valid(row["Check Out At"]) else ""
    if valid(row.get("Late (Minutes)", "X")):
        r["Telat HRIS (mnt)"] = int(float(row["Late (Minutes)"]))
    if str(row["Check In At"])[6:8] == "00":
        cat.append("check-in input manual")

    sh, jad, status, bi = tentukan_shift(row, cfg, shifts)
    r["Shift Aktual"] = f"{hhmm(sh['masuk'])}-{hhmm(sh['pulang'])}"
    r["Status Shift"] = status
    if bi > cfg["review_jika_biaya_di_atas_menit"]:
        cat.append(f"REVIEW: jam absen tidak cocok dengan shift mana pun (selisih {int(bi)} mnt)")

    r["Telat (mnt)"] = max(0, int(ci - sh["masuk"]) - cfg["toleransi_telat_menit"])
    if r["Telat (mnt)"] > cfg["penggajian"]["potongan_telat_ambang_menit"]:
        r["Potongan Telat (Rp)"] = cfg["penggajian"]["potongan_telat_per_kejadian"]

    if co is None or co <= ci:
        cat.append("REVIEW: check-out kosong/tidak valid (cek lupa absen pulang)")
    else:
        kerja = co - ci
        if valid(row.get("Break Start", "X")) and valid(row.get("Break Finish", "X")):
            ks, kf = menit(row["Break Start"]), menit(row["Break Finish"])
            if ks is not None and kf is not None and kf > ks:
                kerja -= kf - ks
        r["Jam Kerja"] = round(kerja / 60, 2)
        if kerja < cfg["min_menit_kerja_valid"]:
            cat.append(f"REVIEW: durasi kerja hanya {int(kerja)} mnt (check-out salah?)")
        else:
            pc = int(sh["pulang"] - co) - cfg["toleransi_pulang_cepat_menit"]
            r["Pulang Cepat (mnt)"] = max(0, pc)
            lembur = int(co - sh["pulang"])
            r["Lembur (mnt)"] = max(0, lembur)
            tarif = aturan_lembur(lembur_rules, row["Branch"], r["Tanggal"], row["NIP"])
            if lembur >= cfg["lembur"]["min_menit"]:
                if tarif is not None:
                    p = cfg["lembur"]["pembulatan_menit"]
                    r["Lembur Dibayar (jam)"] = (lembur // p) * p / 60
                    r["Upah Lembur (Rp)"] = round(r["Lembur Dibayar (jam)"] * tarif)
                    cat.append("lembur dibayar (aturan manual)")
                else:
                    cat.append("lembur terdeteksi, tidak ada aturan lembur -> tidak dibayar")
            if lembur >= 240:
                cat.append("REVIEW: lembur >= 4 jam, pastikan jam pulang benar")

    if valid(row.get("Double Shift Check In At", "X")):
        r["Double Shift"] = "ya"
        cat.append("double shift (shift tambahan)")
    if status.startswith("shift berubah"):
        cat.append(f"shift berubah dari {r['Jadwal HRIS']}")
    return r, cat


def rekap_karyawan(det, df, cfg, izin, master):
    pen = cfg["penggajian"]
    rows = []
    for (nama, nip), g in det.groupby(["Nama", "NIP"], sort=False):
        hadir = g[g["Status"] == "Present"]
        absen = g[g["Status"] == "Absent"]
        n_izin = n_bayar = 0
        if izin is not None:
            iz = izin[(izin["nip"] == nip) | (izin["nama"].str.lower() == nama.lower())]
            tg = set(pd.to_datetime(iz["tanggal"]).dt.date)
            n_izin = int(absen["Tanggal"].isin(tg).sum())
            n_bayar = int(absen[absen["Tanggal"].isin(tg)].merge(
                iz.assign(Tanggal=pd.to_datetime(iz["tanggal"]).dt.date), on="Tanggal")["dibayar"].str.lower().eq("ya").sum())
        absen_potong = len(absen) - n_izin + (n_izin - n_bayar)
        telat = hadir[hadir["Telat (mnt)"] > 0]
        pc = hadir[hadir["Pulang Cepat (mnt)"] > 0]
        r = {"Nama": nama, "NIP": nip, "Cabang": g["Cabang"].iloc[0], "Posisi": g["Posisi"].iloc[0],
             "Hari Terjadwal": len(hadir) + len(absen), "Hari Libur (Off)": int((g["Status"] == "Off").sum()),
             "Hadir": len(hadir), "Absen": len(absen), "Izin/Sakit/Cuti": n_izin,
             "Absen Dipotong": absen_potong,
             "Telat (kali)": len(telat), "Telat (mnt)": int(telat["Telat (mnt)"].sum()),
             "Pulang Cepat (kali)": len(pc), "Pulang Cepat (mnt)": int(pc["Pulang Cepat (mnt)"].sum()),
             "Lembur Dibayar (jam)": float(hadir["Lembur Dibayar (jam)"].sum()),
             "Lembur Tidak Dibayar (jam)": round(float(hadir[hadir["Lembur Dibayar (jam)"] == 0]["Lembur (mnt)"]
                                                     .where(lambda x: x >= cfg["lembur"]["min_menit"], 0).sum()) / 60, 1),
             "Telat > Ambang (kali)": int((hadir["Potongan Telat (Rp)"] > 0).sum()),
             "Potongan Telat (Rp)": int(hadir["Potongan Telat (Rp)"].sum()),
             "Upah Lembur (Rp)": int(hadir["Upah Lembur (Rp)"].sum()),
             "Double Shift": int((hadir["Double Shift"] == "ya").sum()),
             "Shift Berubah (kali)": int(hadir["Status Shift"].str.startswith("shift berubah").sum()),
             "Telat Menurut HRIS (mnt)": int(pd.to_numeric(hadir["Telat HRIS (mnt)"], errors="coerce").fillna(0).sum())}
        r["Selisih vs HRIS (mnt)"] = r["Telat (mnt)"] - r["Telat Menurut HRIS (mnt)"]
        rows.append(r)
    out = pd.DataFrame(rows)

    if master is not None:
        m = master.copy()
        for c in ("gaji_pokok", "tunjangan_tetap", "uang_makan_per_hari"):
            m[c] = pd.to_numeric(m.get(c, 0), errors="coerce").fillna(0)
        out = out.merge(m[["nip", "gaji_pokok", "tunjangan_tetap", "uang_makan_per_hari"]],
                        left_on="NIP", right_on="nip", how="left").drop(columns="nip")
        for c in ("gaji_pokok", "tunjangan_tetap", "uang_makan_per_hari"):
            out[c] = out[c].fillna(0)
        out["Potongan Absen"] = (out["gaji_pokok"] / pen["pembagi_hari_kerja"] * out["Absen Dipotong"]).round()
        out["Potongan Pulang Cepat"] = out["Pulang Cepat (mnt)"] * pen["potongan_pulang_cepat_per_menit"]
        out["Uang Makan"] = out["uang_makan_per_hari"] * out["Hadir"]
        out["Gaji Bersih"] = (out["gaji_pokok"] + out["tunjangan_tetap"] + out["Uang Makan"] + out["Upah Lembur (Rp)"]
                              + out["Double Shift"] * pen["bonus_double_shift_per_shift"]
                              - out["Potongan Absen"] - out["Potongan Telat (Rp)"] - out["Potongan Pulang Cepat"])
    return out


def kumpulkan_review(det, catatan, rekap, df):
    rows = []
    for (i, r), cat in zip(det.iterrows(), catatan):
        review = [c for c in cat if c.startswith("REVIEW")]
        if review:
            rows.append({"Tanggal": r["Tanggal"], "Nama": r["Nama"], "NIP": r["NIP"], "Cabang": r["Cabang"],
                         "Jadwal HRIS": r["Jadwal HRIS"], "Check In": r["Check In"], "Check Out": r["Check Out"],
                         "Alasan": "; ".join(c.replace("REVIEW: ", "") for c in review)})
    for _, r in rekap.iterrows():
        if r["Hari Terjadwal"] and r["Hadir"] == 0:
            rows.append({"Tanggal": "", "Nama": r["Nama"], "NIP": r["NIP"], "Cabang": r["Cabang"], "Jadwal HRIS": "",
                         "Check In": "", "Check Out": "",
                         "Alasan": f"Tidak pernah absen sebulan penuh ({r['Absen']} hari Absent) - resign/belum pakai app/izin?"})
    dup = df.groupby("NIP")["Employee Name"].nunique()
    for nip in dup[dup > 1].index:
        rows.append({"Tanggal": "", "Nama": ", ".join(df[df["NIP"] == nip]["Employee Name"].unique()), "NIP": nip,
                     "Cabang": "", "Jadwal HRIS": "", "Check In": "", "Check Out": "",
                     "Alasan": "NIP dipakai lebih dari satu orang - rapikan di HRIS"})
    for _, r in df[df["NIP"].isin(["0", "", "nan"])][["Employee Name", "NIP", "Branch"]].drop_duplicates().iterrows():
        rows.append({"Tanggal": "", "Nama": r["Employee Name"], "NIP": r["NIP"], "Cabang": r["Branch"],
                     "Jadwal HRIS": "", "Check In": "", "Check Out": "", "Alasan": "NIP kosong/0 - lengkapi di HRIS"})
    return pd.DataFrame(rows)


# ---------- output ----------
def format_sheet(ws, lebar=None):
    head = PatternFill("solid", fgColor="1F4E78")
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = head
        c.alignment = Alignment(wrap_text=True, vertical="center")
    ws.freeze_panes = "D2" if ws.title != "Perlu Review" else "C2"
    ws.auto_filter.ref = ws.dimensions
    for i, col in enumerate(ws.columns, 1):
        w = max(len(str(c.value)) if c.value is not None else 0 for c in col)
        ws.column_dimensions[get_column_letter(i)].width = min(max(10, w + 2), 60)


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("export", help="file .xlsx hasil export HRIS")
    ap.add_argument("-o", "--output", help="path file hasil (default output/Rekap_Absensi_YYYY-MM.xlsx)")
    a = ap.parse_args()

    cfg = json.loads((BASE / "config.json").read_text(encoding="utf-8"))
    shifts = siapkan_shift(cfg)
    df = baca_export(a.export)
    izin, master = baca_opsional("izin_cuti.csv"), baca_opsional("karyawan.csv")
    lembur_rules = baca_opsional("lembur.csv")

    hasil = [hitung_hari(r, cfg, shifts, lembur_rules) for _, r in df.iterrows()]
    det = pd.DataFrame([h[0] for h in hasil])
    catatan = [h[1] for h in hasil]
    det["Catatan"] = ["; ".join(c) for c in catatan]
    det = det.sort_values(["Nama", "Tanggal"]).reset_index(drop=True)
    # urutan catatan harus mengikuti urutan det setelah sort
    catatan = [c.split("; ") if c else [] for c in det["Catatan"]]

    rekap = rekap_karyawan(det, df, cfg, izin, master)
    review = kumpulkan_review(det, catatan, rekap, df)

    ringkas = det[det["Status"] == "Present"].groupby("Cabang").agg(
        Hadir=("Status", "size"), Telat_kali=("Telat (mnt)", lambda s: int((s > 0).sum())),
        Telat_mnt=("Telat (mnt)", "sum"), Lembur_jam=("Lembur Dibayar (jam)", "sum")).reset_index()
    ringkas["Absen"] = ringkas["Cabang"].map(det[det["Status"] == "Absent"].groupby("Cabang").size()).fillna(0).astype(int)

    periode = det["Tanggal"].min().strftime("%Y-%m")
    out = Path(a.output) if a.output else BASE / "output" / f"Rekap_Absensi_{periode}.xlsx"
    out.parent.mkdir(parents=True, exist_ok=True)
    with pd.ExcelWriter(out, engine="openpyxl") as xw:
        rekap.to_excel(xw, sheet_name="Rekap Penggajian", index=False)
        det.to_excel(xw, sheet_name="Detail Harian", index=False)
        review.to_excel(xw, sheet_name="Perlu Review", index=False)
        ringkas.to_excel(xw, sheet_name="Ringkasan Cabang", index=False)
        for ws in xw.book.worksheets:
            format_sheet(ws)
    print(f"Selesai -> {out}")
    print(f"  Karyawan: {len(rekap)} | Hadir: {int(rekap['Hadir'].sum())} | Absen: {int(rekap['Absen'].sum())}"
          f" | Telat (kali): {int(rekap['Telat (kali)'].sum())} ({int(rekap['Telat (mnt)'].sum())} mnt)")
    print(f"  Telat menurut HRIS: {int(rekap['Telat Menurut HRIS (mnt)'].sum())} mnt"
          f" | Shift berubah: {int(rekap['Shift Berubah (kali)'].sum())} hari | Perlu review: {len(review)} baris")


if __name__ == "__main__":
    main()
